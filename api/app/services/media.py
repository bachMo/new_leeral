import io
import zipfile
from dataclasses import dataclass
from enum import StrEnum
from pathlib import PurePath

import docx
import pymupdf

from app.ai.contracts import PageInput
from app.core.errors import AppError, ErrorCode

PDF_RENDER_DPI = 150
MIN_PDF_PAGE_TEXT = 40
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class FileKind(StrEnum):
    IMAGE = "image"
    PDF = "pdf"
    DOCX = "docx"
    AUDIO = "audio"


@dataclass(frozen=True, slots=True)
class DetectedFile:
    kind: FileKind
    mime_type: str
    extension: str


@dataclass(frozen=True, slots=True)
class IncomingFile:
    filename: str
    content: bytes


def safe_filename(filename: str | None, default: str) -> str:
    name = PurePath((filename or "").replace("\\", "/")).name.strip()
    return name[:120] or default


def detect(content: bytes) -> DetectedFile:
    if content.startswith(b"\xff\xd8\xff"):
        return DetectedFile(FileKind.IMAGE, "image/jpeg", "jpg")
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return DetectedFile(FileKind.IMAGE, "image/png", "png")
    if content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        return DetectedFile(FileKind.IMAGE, "image/webp", "webp")
    if content.startswith(b"%PDF"):
        return DetectedFile(FileKind.PDF, "application/pdf", "pdf")
    if content.startswith(b"PK\x03\x04") and _is_docx(content):
        return DetectedFile(FileKind.DOCX, DOCX_MIME, "docx")
    raise AppError(ErrorCode.FILE_TYPE_NOT_SUPPORTED)


def detect_audio(content: bytes) -> DetectedFile:
    if content.startswith(b"OggS"):
        return DetectedFile(FileKind.AUDIO, "audio/ogg", "ogg")
    if content[:4] == b"RIFF" and content[8:12] == b"WAVE":
        return DetectedFile(FileKind.AUDIO, "audio/wav", "wav")
    if content.startswith(b"ID3") or content[:2] in {b"\xff\xfb", b"\xff\xf3", b"\xff\xf2"}:
        return DetectedFile(FileKind.AUDIO, "audio/mpeg", "mp3")
    if content[4:8] == b"ftyp":
        return DetectedFile(FileKind.AUDIO, "audio/mp4", "m4a")
    if content.startswith(b"\x1aE\xdf\xa3"):
        return DetectedFile(FileKind.AUDIO, "audio/webm", "webm")
    if content[:2] in {b"\xff\xf1", b"\xff\xf9"}:
        return DetectedFile(FileKind.AUDIO, "audio/aac", "aac")
    raise AppError(ErrorCode.AUDIO_UNREADABLE)


def _is_docx(content: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            return "word/document.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False


def pdf_page_count(content: bytes) -> int:
    try:
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            return int(pdf.page_count)
    except (RuntimeError, ValueError) as exc:
        raise AppError(ErrorCode.DOCUMENT_UNREADABLE) from exc


def pdf_pages(content: bytes, *, start_position: int, max_pages: int) -> list[PageInput]:
    pages: list[PageInput] = []
    with pymupdf.open(stream=content, filetype="pdf") as pdf:
        for index, page in enumerate(pdf):
            if index >= max_pages:
                break
            text = page.get_text("text").strip()
            image: bytes | None = None
            if len(text) < MIN_PDF_PAGE_TEXT:
                image = page.get_pixmap(dpi=PDF_RENDER_DPI).tobytes("png")
            pages.append(
                PageInput(
                    position=start_position + index,
                    mime_type="image/png" if image else "text/plain",
                    image=image,
                    text=text or None,
                    filename=f"page-{start_position + index + 1}.png",
                )
            )
    return pages


def docx_text(content: bytes) -> str:
    document = docx.Document(io.BytesIO(content))
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def build_pages(files: list[tuple[DetectedFile, bytes, str]], max_pages: int) -> list[PageInput]:
    pages: list[PageInput] = []
    for detected, content, filename in files:
        remaining = max_pages - len(pages)
        if remaining <= 0:
            break
        if detected.kind is FileKind.IMAGE:
            pages.append(
                PageInput(
                    position=len(pages),
                    mime_type=detected.mime_type,
                    image=content,
                    filename=filename,
                )
            )
        elif detected.kind is FileKind.PDF:
            pages.extend(pdf_pages(content, start_position=len(pages), max_pages=remaining))
        else:
            pages.append(
                PageInput(
                    position=len(pages),
                    mime_type="text/plain",
                    text=docx_text(content),
                    filename=filename,
                )
            )
    return pages
